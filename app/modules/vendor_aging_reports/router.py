"""Protected API routes for association Vendor Aging Reports."""

from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.constants import RoleCode
from app.core.dependencies import AuthContext, get_auth_context, require_csrf, require_role
from app.core.responses import ApiResponse, success_response
from app.db.session import get_db_session
from app.db.unit_of_work import UnitOfWork
from app.modules.financial_report_imports import (
    MAX_IMPORT_BYTES,
    combined_errors,
    find_header,
    first_nonempty_label,
    non_negative_decimal,
    typed_date,
    validation_error_message,
)
from app.modules.vendor_aging_reports.schemas import (
    VendorAgingMutationResponse,
    VendorAgingReportRequest,
    VendorAgingReportResponse,
    VendorAgingRowRequest,
)
from app.modules.vendor_aging_reports.service import VendorAgingReportService

VIEWER_ROLE_CODES = (
    RoleCode.SUPER_ADMIN,
    RoleCode.ADMIN,
    RoleCode.ACCOUNTANT,
    RoleCode.HOMEOWNER,
    RoleCode.BOARD_MEMBER,
    RoleCode.COMMITTEE_MEMBER,
)
MANAGER_ROLE_CODES = (
    RoleCode.SUPER_ADMIN,
    RoleCode.ADMIN,
    RoleCode.ACCOUNTANT,
)
VIEWER_DEP = Depends(require_role(*VIEWER_ROLE_CODES))
MANAGER_DEP = Depends(require_role(*MANAGER_ROLE_CODES))
router = APIRouter(prefix="/vendor-aging-reports", tags=["Vendor Aging Reports"])


def parse_import_workbook(
    workbook,
    association_id: str,
    fiscal_period: str,
) -> VendorAgingReportRequest:
    rows = list(workbook.active.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The workbook is empty. Use the Vendor Aging template.")
    aliases = {
        "vendor_name": {"vendor", "vendor name"},
        "category": {"category"},
        "invoice_number": {"invoice", "invoice number"},
        "invoice_date": {"date", "invoice date"},
        "current_due": {"current", "current due"},
        "days_1_30": {"1 30", "1 30 days", "days 1 30"},
        "days_31_60": {"31 60", "31 60 days", "days 31 60"},
        "days_61_90": {"61 90", "61 90 days", "days 61 90"},
        "days_over_90": {"90 days", "over 90", "days over 90"},
    }
    header_index, columns = find_header(
        rows,
        aliases,
        tuple(aliases.keys()),
    )
    errors: list[str] = []
    imported_rows: list[VendorAgingRowRequest] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not row or not any(value not in (None, "") for value in row):
            continue
        if first_nonempty_label(row) in {"total", "totals"}:
            break

        def value_for(field: str) -> object:
            index = columns[field]
            return row[index] if index < len(row) else None

        vendor_name = str(value_for("vendor_name") or "").strip()
        category = str(value_for("category") or "").strip()
        invoice_number = str(value_for("invoice_number") or "").strip()
        if not vendor_name or not category or not invoice_number:
            errors.append(
                f"Row {row_number}: Vendor Name, Category, and Invoice Number are required."
            )
            continue
        try:
            imported_rows.append(
                VendorAgingRowRequest(
                    vendor_name=vendor_name,
                    category=category,
                    invoice_number=invoice_number,
                    invoice_date=typed_date(
                        value_for("invoice_date"), row_number, "Invoice Date"
                    ),
                    current_due=non_negative_decimal(
                        value_for("current_due"), row_number, "Current Due"
                    ),
                    days_1_30=non_negative_decimal(
                        value_for("days_1_30"), row_number, "1-30 Days"
                    ),
                    days_31_60=non_negative_decimal(
                        value_for("days_31_60"), row_number, "31-60 Days"
                    ),
                    days_61_90=non_negative_decimal(
                        value_for("days_61_90"), row_number, "61-90 Days"
                    ),
                    days_over_90=non_negative_decimal(
                        value_for("days_over_90"), row_number, "> 90 Days"
                    ),
                )
            )
        except ValidationError as exc:
            errors.append(validation_error_message(exc, row_number))
        except ValueError as exc:
            errors.append(str(exc))
    if errors:
        raise ValueError(combined_errors(errors))
    if not imported_rows:
        raise ValueError("The workbook contains no vendor invoice rows.")
    try:
        return VendorAgingReportRequest(
            association_id=association_id,
            fiscal_period=fiscal_period,
            rows=imported_rows,
        )
    except ValidationError as exc:
        raise ValueError(validation_error_message(exc)) from exc


def build_export_workbook(response: VendorAgingReportResponse) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendor Aging Report"
    amount_format = '#,##0.00;[Red]-#,##0.00'
    sheet.merge_cells("A1:L1")
    sheet["A1"] = f"{response.association_name} - Vendor Aging Report"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="334155")
    sheet.merge_cells("A2:L2")
    category = response.applied_category or "All categories"
    search = response.applied_search or "No search"
    sheet["A2"] = (
        f"Fiscal period: {response.fiscal_period} | Category: {category} | Search: {search}"
    )
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="64748B")

    summaries = (
        ("Vendors", response.summary.vendor_count),
        ("Invoices", response.summary.invoice_count),
        ("Total outstanding", response.summary.total_outstanding),
        ("Overdue", response.summary.overdue_amount),
        ("Current", response.summary.current_amount),
        ("Overdue %", response.summary.overdue_percentage / 100),
    )
    for column, (label, value) in enumerate(summaries, start=1):
        sheet.cell(row=4, column=column, value=label).font = Font(bold=True, color="475569")
        cell = sheet.cell(row=5, column=column, value=float(value))
        if column in {3, 4, 5}:
            cell.number_format = amount_format
        elif column == 6:
            cell.number_format = "0.00%"

    headers = [
        "#",
        "Vendor Name",
        "Category",
        "Invoice Number",
        "Invoice Date",
        "Current Due (₹)",
        "1-30 Days (₹)",
        "31-60 Days (₹)",
        "61-90 Days (₹)",
        "> 90 Days (₹)",
        "Total Outstanding (₹)",
        "Status",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=7, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="right" if 6 <= column <= 11 else "left")

    row_number = 8
    for index, item in enumerate(response.rows, start=1):
        values = [
            index,
            item.vendor_name,
            item.category,
            item.invoice_number,
            item.invoice_date,
            float(item.current_due),
            float(item.days_1_30),
            float(item.days_31_60),
            float(item.days_61_90),
            float(item.days_over_90),
            float(item.total_outstanding),
            item.status.value.replace("_", " ").title(),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            if 6 <= column <= 11:
                cell.number_format = amount_format
                cell.alignment = Alignment(horizontal="right")
        row_number += 1

    sheet.cell(row=row_number, column=2, value="TOTALS").font = Font(bold=True)
    totals = [
        sum((item.current_due for item in response.rows), start=0),
        sum((item.days_1_30 for item in response.rows), start=0),
        sum((item.days_31_60 for item in response.rows), start=0),
        sum((item.days_61_90 for item in response.rows), start=0),
        sum((item.days_over_90 for item in response.rows), start=0),
        response.summary.total_outstanding,
    ]
    for column, value in enumerate(totals, start=6):
        cell = sheet.cell(row=row_number, column=column, value=float(value))
        cell.number_format = amount_format
        cell.font = Font(bold=True, color="1D4ED8")

    widths = [6, 28, 23, 18, 15, 18, 17, 17, 17, 17, 22, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A8"
    sheet.sheet_view.showGridLines = False
    sheet.print_area = f"A1:L{row_number}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return workbook


def workbook_response(workbook: Workbook, filename: str) -> StreamingResponse:
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/template", dependencies=[MANAGER_DEP])
async def vendor_aging_report_template() -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendor Aging Report"
    headers = [
        "Vendor Name",
        "Category",
        "Invoice Number",
        "Invoice Date",
        "Current Due",
        "1-30 Days",
        "31-60 Days",
        "61-90 Days",
        "> 90 Days",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for column, width in enumerate((28, 22, 18, 16, 18, 17, 17, 17, 17), start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    return workbook_response(workbook, "vendor_aging_report_import_template.xlsx")


@router.post(
    "/import",
    response_model=ApiResponse[VendorAgingMutationResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP],
)
async def import_vendor_aging_report(
    association_id: str,
    fiscal_period: str,
    file: Annotated[UploadFile, File()],
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            "Select a Vendor Aging workbook in .xlsx format",
        )
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            "The Vendor Aging workbook must be 5 MB or smaller",
        )
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Unable to read the Excel workbook",
        ) from exc
    try:
        payload = parse_import_workbook(workbook, association_id, fiscal_period)
    except ValueError as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            str(exc),
        ) from exc
    async with UnitOfWork(session):
        report = await VendorAgingReportService(session).create(
            context.account,
            payload,
        )
    return success_response(
        VendorAgingMutationResponse(
            id=report.id,
            imported_items=len(payload.rows),
        ),
        "Vendor Aging Report imported successfully",
    )


@router.get("/periods", response_model=ApiResponse[list[str]], dependencies=[VIEWER_DEP])
async def list_vendor_aging_periods(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    periods = await VendorAgingReportService(session).periods(
        context.account,
        association_id,
    )
    return success_response(periods)


@router.get("/export", dependencies=[VIEWER_DEP])
async def export_vendor_aging_report(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    fiscal_period: str | None = None,
    category: str | None = None,
    search: str | None = None,
) -> StreamingResponse:
    service = VendorAgingReportService(session)
    report = await service.get(context.account, association_id, fiscal_period)
    if not report:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Vendor Aging Report not found")
    response = await service.serialize(report, category, search)
    return workbook_response(
        build_export_workbook(response),
        f"vendor_aging_report_{report.id}.xlsx",
    )


@router.get(
    "",
    response_model=ApiResponse[VendorAgingReportResponse | None],
    dependencies=[VIEWER_DEP],
)
async def get_vendor_aging_report(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    fiscal_period: str | None = None,
    category: str | None = None,
    search: str | None = None,
) -> dict:
    service = VendorAgingReportService(session)
    report = await service.get(context.account, association_id, fiscal_period)
    if not report:
        return success_response(None)
    return success_response(await service.serialize(report, category, search))


@router.post(
    "",
    response_model=ApiResponse[VendorAgingMutationResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP],
)
async def create_vendor_aging_report(
    payload: VendorAgingReportRequest,
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    async with UnitOfWork(session):
        report = await VendorAgingReportService(session).create(context.account, payload)
    return success_response(
        VendorAgingMutationResponse(id=report.id),
        "Vendor Aging Report created successfully",
    )


@router.put(
    "/{report_id}",
    response_model=ApiResponse[VendorAgingMutationResponse],
    dependencies=[MANAGER_DEP],
)
async def update_vendor_aging_report(
    report_id: str,
    payload: VendorAgingReportRequest,
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    async with UnitOfWork(session):
        report = await VendorAgingReportService(session).update(
            context.account,
            report_id,
            payload,
        )
    return success_response(
        VendorAgingMutationResponse(id=report.id),
        "Vendor Aging Report updated successfully",
    )
