"""Protected API routes for association Delinquency Reports."""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.constants import RoleCode
from app.core.dependencies import AuthContext, get_auth_context, require_csrf, require_role
from app.core.responses import ApiResponse, success_response
from app.db.session import get_db_session
from app.db.unit_of_work import UnitOfWork
from app.modules.financial_report_imports import (
    combined_errors,
    display_value,
    find_header,
    validation_error_message,
)
from app.modules.delinquency_reports.schemas import (
    DelinquencyMutationResponse,
    DelinquencyReportRequest,
    DelinquencyReportResponse,
    DelinquencyRowRequest,
)
from app.modules.delinquency_reports.service import DelinquencyReportService

VIEWER_ROLE_CODES = (
    RoleCode.SUPER_ADMIN,
    RoleCode.ADMIN,
    RoleCode.ACCOUNTANT,
    RoleCode.HOMEOWNER,
    RoleCode.BOARD_MEMBER,
    RoleCode.COMMITTEE_MEMBER,
)
MANAGER_ROLE_CODES = (
    RoleCode.SUPER_ADMIN,
    RoleCode.ADMIN,
    RoleCode.ACCOUNTANT,
)
VIEWER_DEP = Depends(require_role(*VIEWER_ROLE_CODES))
MANAGER_DEP = Depends(require_role(*MANAGER_ROLE_CODES))
router = APIRouter(prefix="/delinquency-reports", tags=["Delinquency Reports"])
MAX_IMPORT_BYTES = 5 * 1024 * 1024
CHARGE_CATEGORY_VALUES = [
    "Maintenance",
    "Late Fee",
    "Special Assessment",
    "Utilities",
    "Other",
]


def normalize_cell_label(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


def decimal_cell(value: object, row_number: int, field_name: str) -> Decimal:
    try:
        cleaned = str(value).replace(",", "").replace("₹", "").strip()
        amount = Decimal(cleaned).quantize(Decimal("0.01"))
        if not amount.is_finite() or amount < 0:
            raise InvalidOperation
        return amount
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValueError(
            f"Row {row_number}, {field_name}: expected a non-negative number, "
            f"received {display_value(value)}."
        ) from exc


def integer_cell(value: object, row_number: int, field_name: str) -> int:
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite() or number < 0 or number != number.to_integral_value():
            raise InvalidOperation
        result = int(number)
        if result > 36500:
            raise InvalidOperation
        return result
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValueError(
            f"Row {row_number}, {field_name}: expected a whole number from 0 to "
            f"36500, received {display_value(value)}."
        ) from exc


def date_cell(value: object, row_number: int) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for format_string in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, format_string).date()
        except ValueError:
            continue
    raise ValueError(
        f"Row {row_number}, Last Payment Date: expected a valid Excel date or "
        f"YYYY-MM-DD, received {display_value(value)}."
    )


def find_summary_amount(rows: list[tuple], label: str) -> Decimal:
    normalized_labels = {normalize_cell_label(label)}
    if label == "Total billed amount":
        normalized_labels.add("billed")
    elif label == "Total collected amount":
        normalized_labels.add("collected")
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if normalize_cell_label(value) not in normalized_labels:
                continue
            candidates = []
            if column_index + 1 < len(row):
                candidates.append(row[column_index + 1])
            if row_index + 1 < len(rows) and column_index < len(rows[row_index + 1]):
                candidates.append(rows[row_index + 1][column_index])
            for candidate in candidates:
                if candidate not in (None, ""):
                    try:
                        return decimal_cell(candidate, row_index + 1, label)
                    except ValueError:
                        continue
    raise ValueError(f'The workbook is missing a numeric "{label}" value.')


def parse_import_workbook(workbook, association_id: str, fiscal_period: str) -> DelinquencyReportRequest:
    rows = list(workbook.active.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The workbook is empty. Use the Delinquency Report template.")

    aliases = {
        "member_name": {"member", "member name"},
        "unit_number": {"unit", "unit number"},
        "block": {"block"},
        "charge_category": {"category", "charge category"},
        "outstanding_amount": {"outstanding", "outstanding amount"},
        "days_overdue": {"days overdue"},
        "last_payment_date": {"last payment", "last payment date"},
    }
    header_index, columns = find_header(
        rows,
        aliases,
        (
            "member_name",
            "unit_number",
            "charge_category",
            "outstanding_amount",
            "days_overdue",
        ),
    )

    errors: list[str] = []
    imported_rows: list[DelinquencyRowRequest] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not row or not any(value not in (None, "") for value in row):
            continue
        first_text = normalize_cell_label(next((value for value in row if value not in (None, "")), ""))
        if first_text in {"total", "totals"}:
            break

        def value_for(field: str) -> object:
            index = columns.get(field, -1)
            return row[index] if index >= 0 and index < len(row) else None

        member_name = str(value_for("member_name") or "").strip()
        unit_number = str(value_for("unit_number") or "").strip()
        charge_category = str(value_for("charge_category") or "").strip()
        missing_fields = [
            label
            for label, value in (
                ("Member Name", member_name),
                ("Unit Number", unit_number),
                ("Charge Category", charge_category),
            )
            if not value
        ]
        if missing_fields:
            field_list = ", ".join(missing_fields)
            verb = "is" if len(missing_fields) == 1 else "are"
            errors.append(
                f"Row {row_number}: {field_list} {verb} required."
            )
            continue
        if charge_category not in CHARGE_CATEGORY_VALUES:
            allowed_categories = ", ".join(CHARGE_CATEGORY_VALUES)
            errors.append(
                f"Row {row_number}: Charge Category must be one of {allowed_categories}; "
                f"received {display_value(charge_category)}."
            )
            continue
        try:
            imported_rows.append(
                DelinquencyRowRequest(
                    member_name=member_name,
                    unit_number=unit_number,
                    block=str(value_for("block") or "").strip() or None,
                    charge_category=charge_category,
                    outstanding_amount=decimal_cell(
                        value_for("outstanding_amount"),
                        row_number,
                        "Outstanding Amount",
                    ),
                    days_overdue=integer_cell(
                        value_for("days_overdue"),
                        row_number,
                        "Days Overdue",
                    ),
                    last_payment_date=date_cell(
                        value_for("last_payment_date"),
                        row_number,
                    ),
                )
            )
        except ValidationError as exc:
            errors.append(validation_error_message(exc, row_number))
        except ValueError as exc:
            errors.append(str(exc))

    if errors:
        raise ValueError(combined_errors(errors))
    if not imported_rows:
        raise ValueError("The workbook contains no delinquency rows.")

    try:
        return DelinquencyReportRequest(
            association_id=association_id,
            fiscal_period=fiscal_period,
            total_billed_amount=find_summary_amount(rows, "Total billed amount"),
            total_collected_amount=find_summary_amount(rows, "Total collected amount"),
            rows=imported_rows,
        )
    except ValidationError as exc:
        raise ValueError(validation_error_message(exc)) from exc


def build_export_workbook(response: DelinquencyReportResponse) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Delinquency Report"
    amount_format = '#,##0.00;[Red]-#,##0.00'
    sheet.merge_cells("A1:I1")
    sheet["A1"] = f"{response.association_name} - Delinquency Report"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="334155")
    sheet.merge_cells("A2:I2")
    block = response.applied_block or "All blocks"
    search = response.applied_search or "No search"
    sheet["A2"] = (
        f"Fiscal period: {response.fiscal_period} | Block: {block} | Search: {search}"
    )
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="64748B")

    summaries = (
        ("Total outstanding", response.summary.total_outstanding, amount_format),
        ("Members", response.summary.delinquent_member_count, None),
        ("Average days", response.summary.average_days_overdue, "0.00"),
        ("High risk", response.summary.high_risk_count, None),
        ("Billed", response.summary.total_billed_amount, amount_format),
        ("Collected", response.summary.total_collected_amount, amount_format),
        ("Collection rate", response.summary.collection_rate_percentage / 100, "0.00%"),
    )
    for column, (label, value, number_format) in enumerate(summaries, start=1):
        sheet.cell(row=4, column=column, value=label).font = Font(
            bold=True, color="475569"
        )
        cell = sheet.cell(row=5, column=column, value=float(value))
        if number_format:
            cell.number_format = number_format

    headers = [
        "#",
        "Member Name",
        "Unit Number",
        "Block",
        "Charge Category",
        "Outstanding Amount (₹)",
        "Days Overdue",
        "Last Payment Date",
        "Status",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=7, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="DC2626")
        cell.alignment = Alignment(horizontal="right" if column == 6 else "left")

    row_number = 8
    for index, item in enumerate(response.rows, start=1):
        values = [
            index,
            item.member_name,
            item.unit_number,
            item.block or "",
            item.charge_category,
            float(item.outstanding_amount),
            item.days_overdue,
            item.last_payment_date,
            item.status.value.title(),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            if column == 6:
                cell.number_format = amount_format
                cell.alignment = Alignment(horizontal="right")
        row_number += 1

    sheet.cell(row=row_number, column=2, value="TOTALS").font = Font(bold=True)
    total_cell = sheet.cell(
        row=row_number,
        column=6,
        value=float(response.summary.total_outstanding),
    )
    total_cell.number_format = amount_format
    total_cell.font = Font(bold=True, color="B91C1C")

    widths = [6, 28, 16, 15, 24, 24, 16, 20, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A8"
    sheet.sheet_view.showGridLines = False
    sheet.print_area = f"A1:I{row_number}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return workbook


def workbook_response(workbook: Workbook, filename: str) -> StreamingResponse:
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def build_import_template_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Delinquency Report"
    sheet["A1"] = "Total billed amount"
    sheet["A2"] = "Total collected amount"
    sheet.merge_cells("A3:G3")
    sheet["A3"] = (
        "Charge Category options: " + ", ".join(CHARGE_CATEGORY_VALUES)
    )
    sheet["A3"].font = Font(italic=True, color="475569")
    sheet["A3"].alignment = Alignment(horizontal="left")
    for row in (1, 2):
        sheet.cell(row=row, column=1).font = Font(bold=True, color="475569")
        sheet.cell(row=row, column=2).fill = PatternFill("solid", fgColor="FEF3C7")
        sheet.cell(row=row, column=2).number_format = '#,##0.00;[Red]-#,##0.00'

    headers = [
        "Member Name",
        "Unit Number",
        "Block",
        "Charge Category",
        "Outstanding Amount",
        "Days Overdue",
        "Last Payment Date",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="DC2626")

    allowed_values = workbook.create_sheet("Allowed Values")
    allowed_values.append(["Charge Category"])
    for category in CHARGE_CATEGORY_VALUES:
        allowed_values.append([category])
    allowed_values.sheet_state = "hidden"

    category_validation = DataValidation(
        type="list",
        formula1="'Allowed Values'!$A$2:$A$6",
        allow_blank=False,
    )
    category_validation.error = "Choose a Charge Category from the list."
    category_validation.errorTitle = "Charge Category is required"
    sheet.add_data_validation(category_validation)
    category_validation.add("D5:D1004")

    required_field_fill = PatternFill("solid", fgColor="FEE2E2")
    for column in ("A", "B", "D"):
        sheet.conditional_formatting.add(
            f"{column}5:{column}1004",
            FormulaRule(
                formula=[f'AND(COUNTA($A5:$G5)>0,{column}5="")'],
                fill=required_field_fill,
            ),
        )
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 24
    sheet.column_dimensions["E"].width = 22
    sheet.column_dimensions["F"].width = 16
    sheet.column_dimensions["G"].width = 20
    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False
    return workbook


@router.get("/template", dependencies=[MANAGER_DEP])
async def delinquency_report_template() -> StreamingResponse:
    return workbook_response(
        build_import_template_workbook(),
        "delinquency_report_import_template.xlsx",
    )


@router.post(
    "/import",
    response_model=ApiResponse[DelinquencyMutationResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP],
)
async def import_delinquency_report(
    association_id: str,
    fiscal_period: str,
    file: Annotated[UploadFile, File()],
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            "Select a Delinquency Report workbook in .xlsx format",
        )
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            "The Delinquency Report workbook must be 5 MB or smaller",
        )
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Unable to read the Excel workbook",
        ) from exc
    try:
        payload = parse_import_workbook(workbook, association_id, fiscal_period)
    except ValueError as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            str(exc),
        ) from exc

    async with UnitOfWork(session):
        report = await DelinquencyReportService(session).create(context.account, payload)
    return success_response(
        DelinquencyMutationResponse(
            id=report.id,
            imported_items=len(payload.rows),
        ),
        "Delinquency Report imported successfully",
    )


@router.get("/periods", response_model=ApiResponse[list[str]], dependencies=[VIEWER_DEP])
async def list_delinquency_periods(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    periods = await DelinquencyReportService(session).periods(
        context.account,
        association_id,
    )
    return success_response(periods)


@router.get("/export", dependencies=[MANAGER_DEP])
async def export_delinquency_report(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    fiscal_period: str | None = None,
    block: str | None = None,
    search: str | None = None,
) -> StreamingResponse:
    service = DelinquencyReportService(session)
    report = await service.get(context.account, association_id, fiscal_period)
    if not report:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Delinquency Report not found")
    response = await service.serialize(report, block, search)
    return workbook_response(
        build_export_workbook(response),
        f"delinquency_report_{report.id}.xlsx",
    )


@router.get(
    "",
    response_model=ApiResponse[DelinquencyReportResponse | None],
    dependencies=[VIEWER_DEP],
)
async def get_delinquency_report(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    fiscal_period: str | None = None,
    block: str | None = None,
    search: str | None = None,
) -> dict:
    service = DelinquencyReportService(session)
    report = await service.get(context.account, association_id, fiscal_period)
    if not report:
        return success_response(None)
    return success_response(await service.serialize(report, block, search))


@router.post(
    "",
    response_model=ApiResponse[DelinquencyMutationResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP],
)
async def create_delinquency_report(
    payload: DelinquencyReportRequest,
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    async with UnitOfWork(session):
        report = await DelinquencyReportService(session).create(context.account, payload)
    return success_response(
        DelinquencyMutationResponse(id=report.id),
        "Delinquency Report created successfully",
    )


@router.put(
    "/{report_id}",
    response_model=ApiResponse[DelinquencyMutationResponse],
    dependencies=[MANAGER_DEP],
)
async def update_delinquency_report(
    report_id: str,
    payload: DelinquencyReportRequest,
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    async with UnitOfWork(session):
        report = await DelinquencyReportService(session).update(
            context.account,
            report_id,
            payload,
        )
    return success_response(
        DelinquencyMutationResponse(id=report.id),
        "Delinquency Report updated successfully",
    )
