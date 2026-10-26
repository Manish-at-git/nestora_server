from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.constants import RoleCode
from app.core.dependencies import AuthContext, get_auth_context, require_csrf, require_role
from app.core.responses import ApiResponse, success_response
from app.db.session import get_db_session
from app.db.unit_of_work import UnitOfWork
from app.modules.invoices.schemas import (
    InvoiceListResponse,
    InvoiceMutationResponse,
    InvoiceRequest,
    InvoiceResponse,
    InvoiceStatus,
    InvoiceStatusRequest,
    InvoiceVendorResponse,
)
from app.modules.invoices.service import InvoiceService

MANAGER_ROLE_CODES = (RoleCode.SUPER_ADMIN, RoleCode.ADMIN, RoleCode.ACCOUNTANT)
VIEWER_ROLE_CODES = (
    *MANAGER_ROLE_CODES,
    RoleCode.HOMEOWNER,
    RoleCode.BOARD_MEMBER,
    RoleCode.COMMITTEE_MEMBER,
)
VIEWER_DEP = Depends(require_role(*VIEWER_ROLE_CODES))
MANAGER_DEP = Depends(require_role(*MANAGER_ROLE_CODES))
router = APIRouter(prefix="/invoices", tags=["Invoices"])
EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_export(rows: list[InvoiceResponse]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoices"
    headers = [
        "Invoice Number",
        "Vendor",
        "Issue Date",
        "Due Date",
        "Fiscal Period",
        "Payment Terms",
        "Subtotal",
        "Discount",
        "Tax Type",
        "Tax Rate",
        "Tax",
        "Grand Total",
        "Status",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
    for item in rows:
        sheet.append(
            [
                item.invoice_number,
                item.vendor_name,
                item.issue_date,
                item.due_date,
                item.fiscal_period,
                item.payment_terms,
                float(item.subtotal),
                float(item.discount_amount),
                item.tax_type or "",
                float(item.tax_rate),
                float(item.tax_amount),
                float(item.grand_total),
                item.status.value,
            ]
        )
    for column in "GHIJKL":
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0.00;[Red]-#,##0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M{max(1, sheet.max_row)}"
    widths = [20, 28, 14, 14, 14, 22, 16, 16, 18, 12, 16, 18, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    return workbook


@router.get(
    "/vendors", response_model=ApiResponse[list[InvoiceVendorResponse]], dependencies=[VIEWER_DEP]
)
async def list_invoice_vendors(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
):
    return success_response(
        await InvoiceService(session).eligible_vendors(context.account, association_id)
    )


@router.get("/export", dependencies=[MANAGER_DEP])
async def export_invoices(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    vendor_id: str | None = None,
    status_filter: InvoiceStatus | None = None,
    fiscal_period: str | None = None,
    search: str | None = None,
):
    result = await InvoiceService(session).list(
        context.account, association_id, vendor_id, status_filter, fiscal_period, search, 1, 10000
    )
    stream = BytesIO()
    build_export(result.rows).save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type=EXCEL_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="invoices.xlsx"'},
    )


@router.get("", response_model=ApiResponse[InvoiceListResponse], dependencies=[VIEWER_DEP])
async def list_invoices(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    vendor_id: str | None = None,
    status_filter: InvoiceStatus | None = None,
    fiscal_period: str | None = None,
    search: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
):
    result = await InvoiceService(session).list(
        context.account,
        association_id,
        vendor_id,
        status_filter,
        fiscal_period,
        search,
        page,
        page_size,
    )
    return success_response(result)


@router.get("/{invoice_id}", response_model=ApiResponse[InvoiceResponse], dependencies=[VIEWER_DEP])
async def get_invoice(
    invoice_id: str,
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
):
    invoice = await InvoiceService(session).get(context.account, invoice_id, association_id)
    return success_response(InvoiceService.serialize(invoice))


@router.post(
    "",
    response_model=ApiResponse[InvoiceResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP, Depends(require_csrf)],
)
async def create_invoice(
    payload: InvoiceRequest,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
):
    service = InvoiceService(session)
    async with UnitOfWork(session):
        invoice = await service.create(context.account, payload)
    return success_response(service.serialize(invoice), "Invoice created")


@router.put(
    "/{invoice_id}",
    response_model=ApiResponse[InvoiceResponse],
    dependencies=[MANAGER_DEP, Depends(require_csrf)],
)
async def update_invoice(
    invoice_id: str,
    payload: InvoiceRequest,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
):
    service = InvoiceService(session)
    async with UnitOfWork(session):
        invoice = await service.update(context.account, invoice_id, payload)
    return success_response(service.serialize(invoice), "Invoice updated")


@router.patch(
    "/{invoice_id}/status",
    response_model=ApiResponse[InvoiceResponse],
    dependencies=[MANAGER_DEP, Depends(require_csrf)],
)
async def update_invoice_status(
    invoice_id: str,
    payload: InvoiceStatusRequest,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
):
    service = InvoiceService(session)
    async with UnitOfWork(session):
        invoice = await service.update_status(context.account, invoice_id, payload)
    return success_response(service.serialize(invoice), "Invoice status updated")


@router.delete(
    "/{invoice_id}",
    response_model=ApiResponse[InvoiceMutationResponse],
    dependencies=[MANAGER_DEP, Depends(require_csrf)],
)
async def delete_invoice(
    invoice_id: str,
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
):
    service = InvoiceService(session)
    async with UnitOfWork(session):
        invoice = await service.delete(context.account, invoice_id, association_id)
    return success_response(InvoiceMutationResponse(id=invoice.id), "Invoice deleted")
