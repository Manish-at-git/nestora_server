"""Association directory, onboarding, and subscription business rules."""

import uuid
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from fastapi import HTTPException, status
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.access_codes import generate_access_code
from app.core.contact_normalization import normalize_email, normalize_phone
from app.core.constants import AccountStatus, RoleCode
from app.core.security import generate_secret, hash_password
from app.modules.auth.models import Account
from app.modules.employees.models import Employee
from app.modules.entities.models import Entity
from app.modules.iam.models import Role
from app.modules.locations.models import City, Country, Region
from app.modules.associations.messages import AssociationMessage
from app.modules.associations.models import Association
from app.modules.associations.repository import AssociationRepository
from app.modules.associations.schemas import (
    AssociationOnboardResponse,
    AssociationSettingsUpdateRequest,
    AssociationSubscriptionRequest,
)
from app.modules.subscriptions.models import SubscriptionPlan
from app.modules.users.models import UserCode, UserDetail


class AssociationService:
    def __init__(self, session: AsyncSession) -> None:
        self.repository = AssociationRepository(session)
        self.registration_deliveries: list[dict[str, str]] = []

    @staticmethod
    def workbook_metrics(workbook_bytes: bytes) -> dict[str, int]:
        """Calculate editable onboarding defaults from the Unit Details sheet."""
        try:
            workbook = load_workbook(
                filename=__import__("io").BytesIO(workbook_bytes),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, AssociationMessage.INVALID_WORKBOOK) from exc

        required = {"Association Details", "Unit Details", "Homeowner Details"}
        if not required.issubset(workbook.sheetnames):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, AssociationMessage.WORKBOOK_SHEETS_REQUIRED)

        values = list(workbook["Unit Details"].values)
        if not values:
            return {"num_blocks": 0, "floors_per_block": 0, "units_per_floor": 0}
        headers = [str(value).strip() if value is not None else "" for value in values[0]]

        def clean(value: object) -> str:
            return "" if value is None else str(value).strip()

        floors_by_block: dict[str, set[str]] = {}
        units_by_floor: dict[tuple[str, str], set[str]] = {}
        for row_values in values[1:]:
            if not any(value is not None for value in row_values):
                continue
            row = dict(zip(headers, row_values))
            block, floor, unit = clean(row.get("Block Name")), clean(row.get("Floor")), clean(row.get("Unit Number"))
            if not block or not unit:
                continue
            floors_by_block.setdefault(block, set()).add(floor)
            units_by_floor.setdefault((block, floor), set()).add(unit)

        return {
            "num_blocks": len(floors_by_block),
            "floors_per_block": max((len(floors) for floors in floors_by_block.values()), default=0),
            "units_per_floor": max((len(units) for units in units_by_floor.values()), default=0),
        }

    async def list(self, account_id: str, role_code: str):
        associations = await self.repository.list(account_id, role_code)
        for association in associations:
            association.unit_count = await self.repository.unit_count(association.id)
            association.allowed_features = await self.repository.allowed_features(association.current_plan_id)
        return associations

    async def stats(self, association_id: str) -> dict:
        association = await self.repository.get(association_id)
        if association is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, AssociationMessage.NOT_FOUND)
        result = await self.repository.stats(association_id)
        return {"name": association.name, "contract_url": association.contract_url, **result}

    async def settings(self, association_id: str, account_id: str, role_code: str) -> dict:
        await self._require_settings_access(association_id, account_id, role_code)
        settings = await self.repository.settings(association_id)
        if settings is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, AssociationMessage.NOT_FOUND)
        return settings

    async def update_settings(
        self,
        association_id: str,
        payload: AssociationSettingsUpdateRequest,
        account_id: str,
        role_code: str,
    ) -> None:
        await self._require_settings_access(association_id, account_id, role_code)
        updated = await self.repository.update_settings(
            association_id=association_id,
            end_date=payload.end_date,
            assessment_rules=(
                payload.assessment_rules.model_dump() if payload.assessment_rules else None
            ),
            fine_rules=(
                [fine_rule.model_dump(exclude={"id"}) for fine_rule in payload.fine_rules]
                if payload.fine_rules is not None
                else None
            ),
            update_end_date="end_date" in payload.model_fields_set,
        )
        if not updated:
            raise HTTPException(status.HTTP_404_NOT_FOUND, AssociationMessage.NOT_FOUND)

    async def _require_settings_access(
        self, association_id: str, account_id: str, role_code: str
    ) -> None:
        if role_code == RoleCode.SUPER_ADMIN:
            return
        if role_code == RoleCode.ADMIN and association_id in await self.repository.association_ids_for_admin(account_id):
            return
        raise HTTPException(status.HTTP_403_FORBIDDEN, AssociationMessage.FORBIDDEN)

    async def update_subscription(self, association_id: str, payload: AssociationSubscriptionRequest) -> None:
        association = await self.repository.get(association_id)
        if association is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, AssociationMessage.NOT_FOUND)
        if payload.plan_id and await self.repository.session.scalar(
            select(SubscriptionPlan.id).where(
                SubscriptionPlan.id == payload.plan_id,
                SubscriptionPlan.is_active.is_(True),
                SubscriptionPlan.is_deleted.is_(False),
            )
        ) is None:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, AssociationMessage.PLAN_NOT_FOUND)
        values = payload.model_dump()
        association.current_plan_id = values.pop("plan_id")
        for key, value in values.items():
            setattr(association, key, value)
        await self.repository.session.flush()

    async def _validate_workbook_contacts(
        self, homeowner_rows: list[dict], tenant_rows: list[tuple[int, dict]]
    ) -> None:
        """Reject duplicate resident contacts before staging onboarding records."""
        existing_emails: dict[str, str] = {}
        existing_phones: dict[str, str] = {}

        users = await self.repository.session.scalars(
            select(UserDetail).where(UserDetail.is_deleted.is_(False))
        )
        for user in users.all():
            email = normalize_email(user.email)
            phone = normalize_phone(user.contact_number)
            if email:
                existing_emails[email] = "an existing user"
            if phone:
                existing_phones[phone] = "an existing user"

        employees = await self.repository.session.scalars(
            select(Employee).where(Employee.is_deleted.is_(False))
        )
        for employee in employees.all():
            email = normalize_email(employee.email)
            phone = normalize_phone(employee.contact_number)
            if email:
                existing_emails[email] = "an existing employee"
            if phone:
                existing_phones[phone] = "an existing employee"

        accounts = await self.repository.session.scalars(select(Account))
        for account in accounts.all():
            email = normalize_email(account.email)
            if email:
                existing_emails[email] = "an existing account"

        seen_emails: dict[str, tuple[str, int]] = {}
        seen_phones: dict[str, tuple[str, int]] = {}

        def validate_contact(
            source: str, row_number: int, email_value: object, phone_value: object
        ) -> None:
            email = normalize_email(email_value)
            phone = normalize_phone(phone_value)
            if email in existing_emails:
                raise HTTPException(
                    status.HTTP_400_BAD_REQUEST,
                    f"Email already exists for {existing_emails[email]}: {email}",
                )
            if phone and phone in existing_phones:
                raise HTTPException(
                    status.HTTP_400_BAD_REQUEST,
                    f"Phone number already exists for {existing_phones[phone]}: {phone}",
                )
            if email in seen_emails:
                previous_source, previous_row = seen_emails[email]
                raise HTTPException(
                    status.HTTP_400_BAD_REQUEST,
                    f"Duplicate email found: {email}"
                )
            if phone and phone in seen_phones:
                previous_source, previous_row = seen_phones[phone]
                raise HTTPException(
                    status.HTTP_400_BAD_REQUEST,
                    f"Duplicate phone number found in {source} and {previous_source}: "
                    f"{phone} (rows {previous_row} and {row_number})",
                )
            if email:
                seen_emails[email] = (source, row_number)
            if phone:
                seen_phones[phone] = (source, row_number)

        for row_number, row in enumerate(homeowner_rows, start=2):
            validate_contact(
                "Homeowner Details",
                row_number,
                row.get("Email"),
                row.get("Phone Number"),
            )
        for row_number, row in tenant_rows:
            validate_contact(
                "Homeowner Details tenant fields",
                row_number,
                row.get("Tenant Email Id"),
                row.get("Tenant Contact Number"),
            )

    async def onboard(
        self,
        workbook_bytes: bytes,
        entity_id: str,
        plan_id: str | None = None,
        contract_url: str | None = None,
    ) -> AssociationOnboardResponse:
        """Import one workbook as one transaction, including all resident records."""
        entity = await self.repository.session.scalar(
            select(Entity).where(Entity.id == entity_id, Entity.is_deleted.is_(False))
        )
        if entity is None:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, AssociationMessage.ENTITY_NOT_FOUND)
        if plan_id and await self.repository.session.scalar(
            select(SubscriptionPlan.id).where(
                SubscriptionPlan.id == plan_id, SubscriptionPlan.is_deleted.is_(False)
            )
        ) is None:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, AssociationMessage.PLAN_NOT_FOUND)

        try:
            workbook = load_workbook(filename=__import__("io").BytesIO(workbook_bytes), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, AssociationMessage.INVALID_WORKBOOK) from exc
        if not {"Association Details", "Unit Details", "Homeowner Details"}.issubset(workbook.sheetnames):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, AssociationMessage.WORKBOOK_SHEETS_REQUIRED)

        def rows(sheet_name: str) -> list[dict]:
            sheet = workbook[sheet_name]
            values = list(sheet.values)
            if not values:
                return []
            headers = [str(value).strip() if value is not None else "" for value in values[0]]
            return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]

        association_rows = rows("Association Details")
        if not association_rows:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, AssociationMessage.ASSOCIATION_SHEET_EMPTY)
        association_data = association_rows[0]

        def clean(value: object) -> str:
            if value is None:
                return ""
            return str(value).strip()

        homeowner_rows = rows("Homeowner Details")
        tenant_rows = [
            (row_number, row)
            for row_number, row in enumerate(homeowner_rows, start=2)
            if clean(row.get("Rented")).lower() == "yes" and clean(row.get("Tenant Email Id"))
        ]
        await self._validate_workbook_contacts(homeowner_rows, tenant_rows)

        association_name = clean(association_data.get("Association Name")) or "Unknown"
        location = await self.repository.session.execute(
            select(City)
            .join(Region, City.region_id == Region.id)
            .join(Country, Region.country_id == Country.id)
            .where(
                func.lower(City.name) == clean(association_data.get("City")).lower(),
                func.lower(Region.name) == clean(association_data.get("State")).lower(),
                func.lower(Country.name) == clean(association_data.get("Country")).lower(),
            )
        )
        city = location.scalar_one_or_none()
        if city is None:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                "Country, State, and City must match a seeded location.",
            )
        association_code = association_name.upper().replace(" ", "")[:4].ljust(4, "0")
        association_id = str(uuid.uuid4())
        association = Association(
            id=association_id,
            name=association_name,
            association_code=association_code,
            entity_id=entity_id,
            address_line_1=clean(association_data.get("Address 1")) or None,
            address_line_2=clean(association_data.get("Address 2")) or None,
            city_id=city.id,
            pincode=clean(association_data.get("Pin Code")) or None,
            url=clean(association_data.get("Association URL")) or None,
            contract_url=contract_url,
            current_plan_id=plan_id,
            subscription_status="Active" if plan_id else "Trial",
            subscription_start=datetime.utcnow().date() if plan_id else None,
            is_deleted=False,
        )
        self.repository.session.add(association)
        await self.repository.session.flush()

        role_rows = await self.repository.session.scalars(
            select(Role).where(Role.code.in_([RoleCode.HOMEOWNER, RoleCode.TENANT]), Role.is_deleted.is_(False))
        )
        roles = {role.code: role.id for role in role_rows.all()}
        if RoleCode.HOMEOWNER not in roles or RoleCode.TENANT not in roles:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Homeowner and tenant roles must be seeded")

        unit_map: dict[tuple[str, str], str] = {}
        units_created = 0
        for row in rows("Unit Details"):
            block_name, unit_number = clean(row.get("Block Name")), clean(row.get("Unit Number"))
            if not block_name or not unit_number:
                continue
            block_id = unit_map.get((block_name, "__block__"))
            if block_id is None:
                block_id = str(uuid.uuid4())
                await self.repository.session.execute(
                    text("INSERT INTO blocks (id, association_id, name) VALUES (:id, :association_id, :name)"),
                    {"id": block_id, "association_id": association_id, "name": block_name},
                )
                unit_map[(block_name, "__block__")] = block_id
            unit_id = str(uuid.uuid4())
            await self.repository.session.execute(
                text("INSERT INTO units (id, block_id, floor, unit_number) VALUES (:id, :block_id, :floor, :unit_number)"),
                {"id": unit_id, "block_id": block_id, "floor": clean(row.get("Floor")) or None, "unit_number": unit_number},
            )
            unit_map[(block_name, unit_number)] = unit_id
            units_created += 1

        homeowner_count = tenant_count = 0
        homeowners_by_name: dict[tuple[str, str], list[tuple[str, str]]] = {}
        for row in homeowner_rows:
            block_name, unit_number = clean(row.get("Block Name")), clean(row.get("Unit Number"))
            if not block_name or not unit_number:
                continue
            unit_id = unit_map.get((block_name, unit_number))
            if unit_id is None:
                raise HTTPException(status.HTTP_400_BAD_REQUEST, f"{AssociationMessage.INVALID_UNIT_REFERENCE}: {block_name}/{unit_number}")
            first_name, last_name = clean(row.get("First Name")), clean(row.get("Last Name"))
            address = f"{block_name}-{unit_number}, {association.address_line_1 or ''}".strip(", ")
            login_code = UserCode(id=str(uuid.uuid4()), login_code=generate_access_code(), status="active")
            self.repository.session.add(login_code)
            homeowner = UserDetail(
                user_id=str(uuid.uuid4()), code_id=login_code.id, name=f"{first_name} {last_name}".strip(),
                first_name=first_name, last_name=last_name, email=normalize_email(row.get("Email")),
                contact_number=normalize_phone(row.get("Phone Number")), unit_id=unit_id, association_id=association_id,
                address=address, role_id=roles[RoleCode.HOMEOWNER], is_deleted=False,
            )
            self.repository.session.add(homeowner)
            if homeowner.email:
                self.registration_deliveries.append(
                    {
                        "email": homeowner.email.lower(),
                        "name": homeowner.name,
                        "registration_code": login_code.login_code,
                        "association_name": association_name,
                    }
                )
            homeowner_name_key = (
                first_name.casefold(),
                last_name.casefold(),
            )
            homeowners_by_name.setdefault(homeowner_name_key, []).append((homeowner.user_id, unit_id))
            homeowner_count += 1
            if clean(row.get("Rented")).lower() == "yes":
                await self.repository.session.execute(
                    text("UPDATE units SET is_rented = 1 WHERE id = :id"), {"id": unit_id}
                )
                tenant_email = normalize_email(row.get("Tenant Email Id"))
                if tenant_email:
                    tenant_code = UserCode(id=str(uuid.uuid4()), login_code=generate_access_code(), status="active")
                    self.repository.session.add(tenant_code)
                    self.repository.session.add(UserDetail(
                        user_id=str(uuid.uuid4()), code_id=tenant_code.id,
                        name=f"{clean(row.get('Tenant First Name'))} {clean(row.get('Tenant Last Name'))}".strip(),
                        first_name=clean(row.get("Tenant First Name")), last_name=clean(row.get("Tenant Last Name")),
                        email=tenant_email, contact_number=normalize_phone(row.get("Tenant Contact Number")),
                        unit_id=unit_id, association_id=association_id, address=address,
                        role_id=roles[RoleCode.TENANT], is_deleted=False,
                    ))
                    tenant_count += 1

        # Flush the ORM-created homeowners/tenants before inserting memberships.
        await self.repository.session.flush()

        if "Board Members" in workbook.sheetnames:
            board_member_rows = rows("Board Members")
        elif "Board & Committee Members" in workbook.sheetnames:
            # Backward-compatible title support; Committee Name is ignored.
            board_member_rows = rows("Board & Committee Members")
        else:
            board_member_rows = []

        board_role_id = await self.repository.session.scalar(
            select(Role.id).where(Role.code == RoleCode.BOARD_MEMBER, Role.is_deleted.is_(False))
        )
        board_members_created = 0
        term_start = date.today()
        term_end = term_start + timedelta(days=365)
        for row in board_member_rows:
            board_member_name_key = (
                clean(row.get("First Name")).casefold(),
                clean(row.get("Last Name")).casefold(),
            )
            matching_homeowners = homeowners_by_name.get(board_member_name_key, [])
            if not matching_homeowners:
                # A board row is eligible only when its name is present in the
                # Homeowner Details sheet. Unit data is retained for identifying
                # the correct homeowner when the same name appears more than once.
                continue

            board_unit_id = unit_map.get((clean(row.get("Block Name")), clean(row.get("Unit Number"))))
            if board_unit_id:
                unit_matches = [
                    user_id for user_id, unit_id in matching_homeowners if unit_id == board_unit_id
                ]
                user_id = unit_matches[0] if unit_matches else (
                    matching_homeowners[0][0] if len(matching_homeowners) == 1 else None
                )
            else:
                user_id = matching_homeowners[0][0] if len(matching_homeowners) == 1 else None
            account = await self.repository.session.scalar(
                select(Account).where(Account.user_id == user_id)
            ) if user_id else None
            if user_id and account is None and board_role_id:
                homeowner = await self.repository.session.scalar(
                    select(UserDetail).where(UserDetail.user_id == user_id)
                )
                if homeowner is not None:
                    # Board membership is assigned during onboarding, before the
                    # resident has completed account setup with their access code.
                    # The inactive account preserves the existing account-based
                    # board-member contract and is activated during registration.
                    account = Account(
                        id=str(uuid.uuid4()),
                        user_id=user_id,
                        email=homeowner.email.lower(),
                        password_hash=hash_password(generate_secret()),
                        role_id=board_role_id,
                        status=AccountStatus.INACTIVE,
                    )
                    self.repository.session.add(account)
                    await self.repository.session.flush()
            account_id = account.id if account else None
            if account_id and board_role_id:
                await self.repository.session.execute(
                    text(
                        "INSERT INTO board_members "
                        "(id, association_id, account_id, term_start_date, term_end_date, status, is_deleted) "
                        "VALUES (:id, :association_id, :account_id, :term_start_date, :term_end_date, 'active', 0)"
                    ),
                    {
                        "id": str(uuid.uuid4()),
                        "association_id": association_id,
                        "account_id": account_id,
                        "term_start_date": term_start,
                        "term_end_date": term_end,
                    },
                )
                await self.repository.session.execute(
                    text("UPDATE accounts SET role_id=:role_id WHERE account_id=:account_id"),
                    {"role_id": board_role_id, "account_id": account_id},
                )
                board_members_created += 1
        return AssociationOnboardResponse(
            association_id=association_id, homeowners_created=homeowner_count,
            tenants_created=tenant_count, units_created=units_created,
            committees_created=0,
        )
