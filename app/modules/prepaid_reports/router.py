"""Protected API routes for association Prepaid Reports."""

from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.constants import RoleCode
from app.core.dependencies import AuthContext, get_auth_context, require_csrf, require_role
from app.core.responses import ApiResponse, success_response
from app.db.session import get_db_session
from app.db.unit_of_work import UnitOfWork
from app.modules.financial_report_imports import (
    MAX_IMPORT_BYTES,
    combined_errors,
    find_header,
    first_nonempty_label,
    non_negative_decimal,
    typed_date,
    validation_error_message,
)
from app.modules.prepaid_reports.schemas import (
    PrepaidReportMutationResponse,
    PrepaidReportRequest,
    PrepaidReportRowRequest,
    PrepaidReportResponse,
)
from app.modules.prepaid_reports.service import PrepaidReportService

VIEWER_ROLE_CODES = (
    RoleCode.SUPER_ADMIN,
    RoleCode.ADMIN,
    RoleCode.ACCOUNTANT,
    RoleCode.HOMEOWNER,
    RoleCode.BOARD_MEMBER,
    RoleCode.COMMITTEE_MEMBER,
)
MANAGER_ROLE_CODES = (
    RoleCode.SUPER_ADMIN,
    RoleCode.ADMIN,
    RoleCode.ACCOUNTANT,
)
VIEWER_DEP = Depends(require_role(*VIEWER_ROLE_CODES))
MANAGER_DEP = Depends(require_role(*MANAGER_ROLE_CODES))
router = APIRouter(prefix="/prepaid-reports", tags=["Prepaid Reports"])


def parse_import_workbook(
    workbook,
    association_id: str,
    fiscal_period: str,
) -> PrepaidReportRequest:
    rows = list(workbook.active.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The workbook is empty. Use the Prepaid Report template.")
    aliases = {
        "member_name": {"member", "member name"},
        "unit_number": {"unit", "unit number"},
        "block": {"block"},
        "category": {"category"},
        "prepaid_amount": {"prepaid", "prepaid amount"},
        "used_amount": {"used", "used amount"},
        "expiry_date": {"expiry", "expiry date", "expiry date status"},
    }
    header_index, columns = find_header(
        rows,
        aliases,
        (
            "member_name",
            "unit_number",
            "category",
            "prepaid_amount",
            "used_amount",
            "expiry_date",
        ),
    )
    errors: list[str] = []
    imported_rows: list[PrepaidReportRowRequest] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not row or not any(value not in (None, "") for value in row):
            continue
        if first_nonempty_label(row) in {"total", "totals"}:
            break

        def value_for(field: str) -> object:
            index = columns.get(field, -1)
            return row[index] if index >= 0 and index < len(row) else None

        member_name = str(value_for("member_name") or "").strip()
        unit_number = str(value_for("unit_number") or "").strip()
        category = str(value_for("category") or "").strip()
        if not member_name or not unit_number or not category:
            errors.append(
                f"Row {row_number}: Member Name, Unit Number, and Category are required."
            )
            continue
        expiry_value = value_for("expiry_date")
        if isinstance(expiry_value, str) and " / " in expiry_value:
            expiry_value = expiry_value.split(" / ", 1)[0]
        try:
            imported_rows.append(
                PrepaidReportRowRequest(
                    member_name=member_name,
                    unit_number=unit_number,
                    block=str(value_for("block") or "").strip() or None,
                    category=category,
                    prepaid_amount=non_negative_decimal(
                        value_for("prepaid_amount"), row_number, "Prepaid Amount"
                    ),
                    used_amount=non_negative_decimal(
                        value_for("used_amount"), row_number, "Used Amount"
                    ),
                    expiry_date=typed_date(expiry_value, row_number, "Expiry Date"),
                )
            )
        except ValidationError as exc:
            errors.append(validation_error_message(exc, row_number))
        except ValueError as exc:
            errors.append(str(exc))
    if errors:
        raise ValueError(combined_errors(errors))
    if not imported_rows:
        raise ValueError("The workbook contains no prepaid rows.")
    try:
        return PrepaidReportRequest(
            association_id=association_id,
            fiscal_period=fiscal_period,
            rows=imported_rows,
        )
    except ValidationError as exc:
        raise ValueError(validation_error_message(exc)) from exc


def build_export_workbook(response: PrepaidReportResponse) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prepaid Report"
    amount_format = '#,##0.00;[Red]-#,##0.00'

    sheet.merge_cells("A1:I1")
    sheet["A1"] = f"{response.association_name} - Prepaid Report"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="334155")
    sheet.merge_cells("A2:I2")
    filter_label = response.applied_block or "All blocks"
    sheet["A2"] = f"Fiscal period: {response.fiscal_period} | Filter: {filter_label}"
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="64748B")

    summaries = (
        ("Total prepaid", response.summary.total_prepaid_amount),
        ("Total used", response.summary.total_used_amount),
        ("Available balance", response.summary.total_available_balance),
        ("Members", response.summary.member_count),
        ("Average prepaid", response.summary.average_prepaid_amount),
        ("Utilization", response.summary.utilization_pct / 100),
    )
    for column, (label, value) in enumerate(summaries, start=1):
        sheet.cell(row=4, column=column, value=label).font = Font(bold=True, color="475569")
        cell = sheet.cell(row=5, column=column, value=float(value))
        if column in {1, 2, 3, 5}:
            cell.number_format = amount_format
        elif column == 6:
            cell.number_format = "0.00%"

    headers = [
        "#",
        "Member Name",
        "Unit Number",
        "Block",
        "Category",
        "Prepaid Amount (₹)",
        "Used Amount (₹)",
        "Available Balance (₹)",
        "Expiry Date / Status",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=7, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="059669")
        cell.alignment = Alignment(horizontal="right" if 6 <= column <= 8 else "left")

    row_number = 8
    for index, item in enumerate(response.rows, start=1):
        values = [
            index,
            item.member_name,
            item.unit_number,
            item.block or "",
            item.category,
            float(item.prepaid_amount),
            float(item.used_amount),
            float(item.available_balance),
            f"{item.expiry_date.isoformat()} / {item.status.value.title()}",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            if 6 <= column <= 8:
                cell.number_format = amount_format
                cell.alignment = Alignment(horizontal="right")
        row_number += 1

    sheet.cell(row=row_number, column=5, value="TOTALS").font = Font(bold=True)
    for column, value in (
        (6, response.summary.total_prepaid_amount),
        (7, response.summary.total_used_amount),
        (8, response.summary.total_available_balance),
    ):
        cell = sheet.cell(row=row_number, column=column, value=float(value))
        cell.number_format = amount_format
        cell.font = Font(bold=True, color="047857")

    widths = [6, 28, 16, 14, 24, 20, 18, 22, 25]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A8"
    sheet.sheet_view.showGridLines = False
    sheet.print_area = f"A1:I{row_number}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return workbook


def workbook_response(workbook: Workbook, filename: str) -> StreamingResponse:
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/template", dependencies=[MANAGER_DEP])
async def prepaid_report_template() -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prepaid Report"
    headers = [
        "Member Name",
        "Unit Number",
        "Block",
        "Category",
        "Prepaid Amount",
        "Used Amount",
        "Expiry Date",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="059669")
    for column, width in enumerate((28, 16, 14, 24, 20, 18, 18), start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    return workbook_response(workbook, "prepaid_report_import_template.xlsx")


@router.post(
    "/import",
    response_model=ApiResponse[PrepaidReportMutationResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP],
)
async def import_prepaid_report(
    association_id: str,
    fiscal_period: str,
    file: Annotated[UploadFile, File()],
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            "Select a Prepaid Report workbook in .xlsx format",
        )
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            "The Prepaid Report workbook must be 5 MB or smaller",
        )
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Unable to read the Excel workbook",
        ) from exc
    try:
        payload = parse_import_workbook(workbook, association_id, fiscal_period)
    except ValueError as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            str(exc),
        ) from exc
    async with UnitOfWork(session):
        report = await PrepaidReportService(session).create(context.account, payload)
    return success_response(
        PrepaidReportMutationResponse(
            id=report.id,
            imported_items=len(payload.rows),
        ),
        "Prepaid Report imported successfully",
    )


@router.get("/periods", response_model=ApiResponse[list[str]], dependencies=[VIEWER_DEP])
async def list_prepaid_report_periods(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    periods = await PrepaidReportService(session).periods(context.account, association_id)
    return success_response(periods)


@router.get("/export", dependencies=[MANAGER_DEP])
async def export_prepaid_report(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    fiscal_period: str | None = None,
    block: str | None = None,
) -> StreamingResponse:
    service = PrepaidReportService(session)
    report = await service.get(context.account, association_id, fiscal_period)
    if not report:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Prepaid Report not found")
    response = await service.serialize(report, block)
    workbook = build_export_workbook(response)
    return workbook_response(workbook, f"prepaid_report_{report.id}.xlsx")


@router.get("", response_model=ApiResponse[PrepaidReportResponse | None], dependencies=[VIEWER_DEP])
async def get_prepaid_report(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    fiscal_period: str | None = None,
    block: str | None = None,
) -> dict:
    service = PrepaidReportService(session)
    report = await service.get(context.account, association_id, fiscal_period)
    if not report:
        return success_response(None)
    return success_response(await service.serialize(report, block))


@router.post(
    "",
    response_model=ApiResponse[PrepaidReportMutationResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP],
)
async def create_prepaid_report(
    payload: PrepaidReportRequest,
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    async with UnitOfWork(session):
        report = await PrepaidReportService(session).create(context.account, payload)
    return success_response(
        PrepaidReportMutationResponse(id=report.id),
        "Prepaid Report created successfully",
    )


@router.put(
    "/{report_id}",
    response_model=ApiResponse[PrepaidReportMutationResponse],
    dependencies=[MANAGER_DEP],
)
async def update_prepaid_report(
    report_id: str,
    payload: PrepaidReportRequest,
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    async with UnitOfWork(session):
        report = await PrepaidReportService(session).update(
            context.account,
            report_id,
            payload,
        )
    return success_response(
        PrepaidReportMutationResponse(id=report.id),
        "Prepaid Report updated successfully",
    )
