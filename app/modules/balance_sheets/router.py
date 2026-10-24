"""Protected API routes for association Balance Sheets."""

from datetime import date
from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.constants import RoleCode
from app.core.dependencies import AuthContext, get_auth_context, require_csrf, require_role
from app.core.responses import ApiResponse, success_response
from app.db.session import get_db_session
from app.db.unit_of_work import UnitOfWork
from app.modules.balance_sheets.schemas import (
    BalanceSheetMutationResponse,
    BalanceSheetRequest,
    BalanceSheetResponse,
)
from app.modules.balance_sheets.service import BalanceSheetService

VIEWER_DEP = Depends(
    require_role(
        RoleCode.SUPER_ADMIN,
        RoleCode.ADMIN,
        RoleCode.ACCOUNTANT,
        RoleCode.BOARD_MEMBER,
        RoleCode.COMMITTEE_MEMBER,
        RoleCode.HOMEOWNER,
    )
)
MANAGER_DEP = Depends(
    require_role(
        RoleCode.SUPER_ADMIN,
        RoleCode.ADMIN,
        RoleCode.ACCOUNTANT,
    )
)
router = APIRouter(prefix="/balance-sheets", tags=["Balance Sheets"])

EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_IMPORT_BYTES = 5 * 1024 * 1024
TEMPLATE_HEADERS = [
    "Section",
    "Account Code (Optional)",
    "Account Name",
    "Amount",
    "Health Status (Optional)",
    "Financial Health Note (Optional)",
]
SECTION_VALUES = [
    "current_asset",
    "non_current_asset",
    "current_liability",
    "non_current_liability",
    "equity",
]


def spreadsheet_response(workbook: Workbook, filename: str) -> StreamingResponse:
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type=EXCEL_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def format_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F1"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["C"].width = 40
    sheet.column_dimensions["D"].width = 18
    sheet.column_dimensions["E"].width = 28
    sheet.column_dimensions["F"].width = 55


def build_template_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Balance Sheet"
    sheet.append(TEMPLATE_HEADERS)
    format_sheet(sheet)

    values_sheet = workbook.create_sheet("Allowed Values")
    values_sheet.append(["Valid Section values"])
    for section in SECTION_VALUES:
        values_sheet.append([section])
    values_sheet.sheet_state = "hidden"

    validation = DataValidation(
        type="list",
        formula1="'Allowed Values'!$A$2:$A$6",
        allow_blank=False,
    )
    validation.error = "Choose one of the supported Balance Sheet sections."
    validation.errorTitle = "Invalid section"
    sheet.add_data_validation(validation)
    validation.add("A2:A1001")
    return workbook


def build_export_workbook(response: BalanceSheetResponse) -> Workbook:
    """Build a presentation-ready Balance Sheet workbook for downloads."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Balance Sheet"

    green = "10B981"
    purple = "4F46E5"
    light_green = "ECFDF5"
    slate = "334155"
    border = Border(bottom=Side(style="thin", color="E2E8F0"))
    amount_format = '#,##0.00;[Red]-#,##0.00'

    sheet.merge_cells("A1:D1")
    sheet["A1"] = f"{response.association_name} - Balance Sheet"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=slate)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26

    sheet.merge_cells("A2:D2")
    sheet["A2"] = f"As on {response.as_on_date.isoformat()}"
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="64748B")

    headers = ["Category", "Particulars", "Amount (₹)", "Notes"]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=value)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="059669")
        cell.alignment = Alignment(horizontal="left" if column != 3 else "right")

    items_by_section: dict[str, list] = {section: [] for section in SECTION_VALUES}
    for item in response.items:
        items_by_section[item.section.value].append(item)

    section_labels = {
        "current_asset": "Current Assets",
        "non_current_asset": "Non-Current Assets",
        "current_liability": "Current Liabilities",
        "non_current_liability": "Non-Current Liabilities",
        "equity": "Equity",
    }
    subtotal_labels = {
        "current_asset": "Total Current Assets",
        "non_current_asset": "Total Non-Current Assets",
        "current_liability": "Total Current Liabilities",
        "non_current_liability": "Total Non-Current Liabilities",
        "equity": "Total Equity",
    }
    subtotals = {
        "current_asset": response.totals.current_assets,
        "non_current_asset": response.totals.non_current_assets,
        "current_liability": response.totals.current_liabilities,
        "non_current_liability": response.totals.non_current_liabilities,
        "equity": response.totals.total_equity,
    }

    row = 5

    def append_section(section: str) -> None:
        nonlocal row
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        section_cell = sheet.cell(row=row, column=1, value=section_labels[section])
        section_cell.font = Font(name="Arial", bold=True, color=slate)
        row += 1

        for item in items_by_section[section]:
            sheet.cell(row=row, column=2, value=item.account_name)
            amount_cell = sheet.cell(row=row, column=3, value=float(item.amount))
            amount_cell.number_format = amount_format
            amount_cell.alignment = Alignment(horizontal="right")
            for column in range(1, 5):
                sheet.cell(row=row, column=column).border = border
            row += 1

        sheet.cell(row=row, column=2, value=subtotal_labels[section])
        total_cell = sheet.cell(row=row, column=3, value=float(subtotals[section]))
        total_cell.number_format = amount_format
        total_cell.alignment = Alignment(horizontal="right")
        for column in range(1, 5):
            cell = sheet.cell(row=row, column=column)
            cell.fill = PatternFill("solid", fgColor=light_green)
            cell.font = Font(name="Arial", bold=True, color="047857")
        row += 2

    for section in SECTION_VALUES[:2]:
        append_section(section)

    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    sheet.cell(row=row, column=1, value="TOTAL ASSETS")
    sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
    assets_total = sheet.cell(row=row, column=3, value=float(response.totals.total_assets))
    assets_total.number_format = amount_format
    assets_total.alignment = Alignment(horizontal="right")
    for column in range(1, 5):
        cell = sheet.cell(row=row, column=column)
        cell.fill = PatternFill("solid", fgColor=green)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
    row += 3

    for section in SECTION_VALUES[2:]:
        append_section(section)

    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    sheet.cell(row=row, column=1, value="TOTAL LIABILITIES & EQUITY")
    sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
    liabilities_total = sheet.cell(
        row=row,
        column=3,
        value=float(response.totals.liabilities_and_equity),
    )
    liabilities_total.number_format = amount_format
    liabilities_total.alignment = Alignment(horizontal="right")
    for column in range(1, 5):
        cell = sheet.cell(row=row, column=column)
        cell.fill = PatternFill("solid", fgColor=purple)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 38
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 30
    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False
    sheet.print_area = f"A1:D{row}"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return workbook


@router.get(
    "/dates",
    response_model=ApiResponse[list[date]],
    dependencies=[VIEWER_DEP],
)
async def list_balance_sheet_dates(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    dates = await BalanceSheetService(session).dates(context.account, association_id)
    return success_response(dates)


@router.get("/template", dependencies=[MANAGER_DEP])
async def download_balance_sheet_template() -> StreamingResponse:
    return spreadsheet_response(
        build_template_workbook(),
        "balance_sheet_import_template.xlsx",
    )


@router.get("/export", dependencies=[VIEWER_DEP])
async def export_balance_sheet(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    as_on_date: date | None = None,
) -> StreamingResponse:
    service = BalanceSheetService(session)
    balance_sheet = await service.get(context.account, association_id, as_on_date)
    if not balance_sheet:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Balance Sheet not found")
    response = await service.serialize(balance_sheet)

    workbook = build_export_workbook(response)

    return spreadsheet_response(
        workbook,
        f"balance_sheet_{balance_sheet.as_on_date.isoformat()}.xlsx",
    )


@router.post(
    "/import",
    response_model=ApiResponse[BalanceSheetMutationResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP],
)
async def import_balance_sheet(
    association_id: str,
    as_on_date: date,
    file: Annotated[UploadFile, File()],
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            "Balance Sheet import requires an .xlsx file",
        )

    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            "Balance Sheet spreadsheet must be 5 MB or smaller",
        )
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Unable to read the Excel workbook",
        ) from exc

    service = BalanceSheetService(session)
    async with UnitOfWork(session):
        balance_sheet = await service.import_workbook(
            context.account,
            association_id,
            as_on_date,
            workbook,
        )
    return success_response(
        BalanceSheetMutationResponse(
            id=balance_sheet.id,
            imported_items=len(balance_sheet.items),
        ),
        "Balance Sheet imported successfully",
    )


@router.get(
    "",
    response_model=ApiResponse[BalanceSheetResponse | None],
    dependencies=[VIEWER_DEP],
)
async def get_balance_sheet(
    association_id: str,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
    as_on_date: date | None = None,
) -> dict:
    service = BalanceSheetService(session)
    balance_sheet = await service.get(context.account, association_id, as_on_date)
    serialized = await service.serialize(balance_sheet) if balance_sheet else None
    return success_response(serialized)


@router.post(
    "",
    response_model=ApiResponse[BalanceSheetMutationResponse],
    status_code=status.HTTP_201_CREATED,
    dependencies=[MANAGER_DEP],
)
async def create_balance_sheet(
    payload: BalanceSheetRequest,
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    async with UnitOfWork(session):
        balance_sheet = await BalanceSheetService(session).create(
            context.account,
            payload,
        )
    return success_response(
        BalanceSheetMutationResponse(id=balance_sheet.id),
        "Balance Sheet created successfully",
    )


@router.put(
    "/{balance_sheet_id}",
    response_model=ApiResponse[BalanceSheetMutationResponse],
    dependencies=[MANAGER_DEP],
)
async def update_balance_sheet(
    balance_sheet_id: str,
    payload: BalanceSheetRequest,
    context: Annotated[AuthContext, Depends(require_csrf)],
    session: Annotated[AsyncSession, Depends(get_db_session)],
) -> dict:
    async with UnitOfWork(session):
        balance_sheet = await BalanceSheetService(session).update(
            context.account,
            balance_sheet_id,
            payload,
        )
    return success_response(
        BalanceSheetMutationResponse(id=balance_sheet.id),
        "Balance Sheet updated successfully",
    )
