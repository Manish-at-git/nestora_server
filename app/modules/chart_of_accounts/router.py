"""Global and association Chart of Accounts routes."""

import csv
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.constants import RoleCode
from app.core.dependencies import AuthContext, get_auth_context, require_csrf, require_role
from app.core.responses import ApiResponse, success_response
from app.db.session import get_db_session
from app.db.unit_of_work import UnitOfWork
from app.modules.chart_of_accounts.schemas import ChartMutationResponse, ChartOfAccountRequest, ChartOfAccountResponse
from app.modules.chart_of_accounts.service import ChartOfAccountsService


MANAGER_DEP = Depends(require_role(RoleCode.SUPER_ADMIN, RoleCode.ADMIN, RoleCode.ACCOUNTANT))
router = APIRouter(tags=["Chart of Accounts"])


@router.get("/admin/global-coa", response_model=ApiResponse[list[ChartOfAccountResponse]], dependencies=[MANAGER_DEP])
async def list_global_coa(session: AsyncSession = Depends(get_db_session)) -> dict:
    rows = await ChartOfAccountsService(session).global_accounts()
    return success_response([ChartOfAccountResponse(**row) for row in rows])


@router.post("/admin/global-coa/upload", response_model=ApiResponse[ChartMutationResponse], status_code=status.HTTP_201_CREATED, dependencies=[MANAGER_DEP])
async def upload_global_coa(file: UploadFile = File(...), context: AuthContext = Depends(require_csrf), session: AsyncSession = Depends(get_db_session)) -> dict:
    content = await file.read()
    try:
        if (file.filename or "").lower().endswith(".csv"):
            workbook = Workbook()
            sheet = workbook.active
            for row in csv.reader(StringIO(content.decode("utf-8-sig"))):
                sheet.append(row)
        else:
            workbook = load_workbook(filename=BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Unable to parse Excel file: {exc}") from exc
    async with UnitOfWork(session):
        count = await ChartOfAccountsService(session).upload_global(workbook, context.account, context.account.role.code)
    return success_response(ChartMutationResponse(message=f"Successfully uploaded {count} global chart accounts"))


@router.get("/accountant/association-coa", response_model=ApiResponse[list[ChartOfAccountResponse]], dependencies=[MANAGER_DEP])
async def list_association_coa(association_id: str = Query(...), context: AuthContext = Depends(get_auth_context), session: AsyncSession = Depends(get_db_session)) -> dict:
    rows = await ChartOfAccountsService(session).association_accounts(context.account, context.account.role.code, association_id)
    return success_response([ChartOfAccountResponse(**row) for row in rows])


@router.get("/accountant/association-coa/status", response_model=ApiResponse[dict[str, int]], dependencies=[MANAGER_DEP])
async def association_coa_status(context: AuthContext = Depends(get_auth_context), session: AsyncSession = Depends(get_db_session)) -> dict:
    return success_response(await ChartOfAccountsService(session).status(context.account, context.account.role.code))


@router.post("/accountant/association-coa/map", response_model=ApiResponse[ChartMutationResponse], dependencies=[MANAGER_DEP])
async def map_association_coa(association_id: str = Query(...), context: AuthContext = Depends(require_csrf), session: AsyncSession = Depends(get_db_session)) -> dict:
    async with UnitOfWork(session):
        count = await ChartOfAccountsService(session).map_global(context.account, context.account.role.code, association_id)
    return success_response(ChartMutationResponse(message=f"Successfully mapped {count} accounts"))


@router.post("/accountant/association-coa", response_model=ApiResponse[ChartMutationResponse], status_code=status.HTTP_201_CREATED, dependencies=[MANAGER_DEP])
async def add_association_coa(payload: ChartOfAccountRequest, association_id: str = Query(...), context: AuthContext = Depends(require_csrf), session: AsyncSession = Depends(get_db_session)) -> dict:
    async with UnitOfWork(session):
        account_id = await ChartOfAccountsService(session).add_association(context.account, context.account.role.code, association_id, payload)
    return success_response(ChartMutationResponse(id=account_id, message="Chart of Account added successfully"))


@router.put("/accountant/association-coa/{account_id}", response_model=ApiResponse[ChartMutationResponse], dependencies=[MANAGER_DEP])
async def update_association_coa(account_id: str, payload: ChartOfAccountRequest, context: AuthContext = Depends(require_csrf), session: AsyncSession = Depends(get_db_session)) -> dict:
    async with UnitOfWork(session):
        await ChartOfAccountsService(session).update_association(context.account, context.account.role.code, account_id, payload)
    return success_response(ChartMutationResponse(message="Chart of Account updated successfully"))


@router.delete("/accountant/association-coa/{account_id}", response_model=ApiResponse[ChartMutationResponse], dependencies=[MANAGER_DEP])
async def delete_association_coa(account_id: str, context: AuthContext = Depends(require_csrf), session: AsyncSession = Depends(get_db_session)) -> dict:
    async with UnitOfWork(session):
        await ChartOfAccountsService(session).delete_association(context.account, context.account.role.code, account_id)
    return success_response(ChartMutationResponse(message="Chart of Account deleted successfully"))


@router.get("/accountant/association-coa/download", dependencies=[MANAGER_DEP])
async def download_association_coa(association_id: str = Query(...), context: AuthContext = Depends(get_auth_context), session: AsyncSession = Depends(get_db_session)):
    rows = await ChartOfAccountsService(session).association_accounts(context.account, context.account.role.code, association_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Chart of Accounts"
    sheet.append(["GL Code", "GL Name", "Structure", "Grouping"])
    for row in rows:
        sheet.append([row["gl_code"], row["gl_name"], row.get("structure"), row.get("grouping")])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="Chart of Accounts.xlsx"'})
